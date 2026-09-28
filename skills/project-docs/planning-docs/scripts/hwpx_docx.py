"""HWPX(OWPML) 문서를 읽어 문서 모델로 바꾸고, 그 모델을 DOCX 로 그린다.

    python hwpx_docx.py <입력.hwpx> <출력.docx>      # 양식 변환

다른 스크립트는 parse_hwpx() 없이 모델을 직접 만들어 render() 만 써도 된다 (proposal_build.py 참고).
지원 범위는 이 프로젝트 양식에 나온 것들이다: 문단(정렬·여백·줄간격·쪽 나눔), 글자(글꼴·크기·굵게·색),
탭(점선 오른쪽 탭), 표(병합·너비·테두리·바탕색·세로정렬), 그림(글자처럼), 개요 번호, 머리말/꼬리말(쪽 번호).
"""
import glob, os, re, sys, zipfile
from lxml import etree
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Emu, RGBColor, Mm

NS = {'hp': 'http://www.hancom.co.kr/hwpml/2011/paragraph', 'hh': 'http://www.hancom.co.kr/hwpml/2011/head',
      'hc': 'http://www.hancom.co.kr/hwpml/2011/core', 'hs': 'http://www.hancom.co.kr/hwpml/2011/section'}
HP, HH, HC = ('{%s}' % NS[k] for k in ('hp', 'hh', 'hc'))
ALIGN = {'LEFT': WD_ALIGN_PARAGRAPH.LEFT, 'CENTER': WD_ALIGN_PARAGRAPH.CENTER, 'RIGHT': WD_ALIGN_PARAGRAPH.RIGHT,
         'JUSTIFY': WD_ALIGN_PARAGRAPH.JUSTIFY, 'DISTRIBUTE': WD_ALIGN_PARAGRAPH.DISTRIBUTE}


def hu2pt(v):            # 1 HWPUNIT = 1/7200 inch = 0.01 pt
    return float(v) / 100.0


def hu2emu(v):
    return int(float(v) * 914400 / 7200)


# ---------------------------------------------------------------- 모델 생성기 (다른 스크립트가 쓴다)
def run(text, font='나눔고딕', size=10, bold=False, color=None, tab=False, field=None):
    return dict(text=text, font=font, size=size, bold=bold, color=color, tab=tab, field=field)


def para(runs=(), align='JUSTIFY', left=0, space_before=0, line=160, page_break=False, tabs=(), border_bottom=None,
         line_width=None, leader=True):
    return dict(type='p', runs=list(runs), align=align, left=left, space_before=space_before, line=line,
                page_break=page_break, tabs=list(tabs), border_bottom=border_bottom, line_width=line_width, leader=leader)


def pic(path, w_pt, h_pt, align='CENTER', page_break=False):
    return dict(type='pic', path=path, w=w_pt, h=h_pt, align=align, page_break=page_break)


def cell(blocks, span=(1, 1), borders=None, fill=None, valign='CENTER', width=0):
    return dict(blocks=blocks, span=span, borders=borders or {}, fill=fill, valign=valign, width=width)


def table(rows, col_widths, align='CENTER', page_break=False):
    return dict(type='table', rows=rows, col_widths=list(col_widths), align=align, page_break=page_break)


BORDER_THIN = {s: ('SOLID', 0.12, '#000000') for s in ('left', 'right', 'top', 'bottom')}
BORDER_NONE = {s: ('NONE', 0, '#000000') for s in ('left', 'right', 'top', 'bottom')}


# ---------------------------------------------------------------- HWPX 파싱
class Hwpx:
    def __init__(self, path):
        self.z = zipfile.ZipFile(path)
        self.head = etree.fromstring(self.z.read('Contents/header.xml'))
        self.sec = etree.fromstring(self.z.read('Contents/section0.xml'))
        self.fonts = {}
        for ff in self.head.iter(HH + 'fontface'):
            for f in ff.findall('hh:font', NS):
                self.fonts[(ff.get('lang'), f.get('id'))] = f.get('face')
        self.charprs = {c.get('id'): c for c in self.head.iter(HH + 'charPr')}
        self.paraprs = {p.get('id'): p for p in self.head.iter(HH + 'paraPr')}
        self.bfs = {b.get('id'): b for b in self.head.iter(HH + 'borderFill')}
        self.numberings = {n.get('id'): n for n in self.head.iter(HH + 'numbering')}
        self.outline = 0
        self.tmpdir = None

    # --- 스타일
    def char(self, cid):
        c = self.charprs.get(cid)
        if c is None:
            return dict(font='나눔고딕', size=10, bold=False, color=None)
        fr = c.find('hh:fontRef', NS)
        face = self.fonts.get(('HANGUL', fr.get('hangul'))) if fr is not None else None
        color = c.get('textColor')
        return dict(font=face or '나눔고딕', size=hu2pt(c.get('height', '1000')), bold=c.find('hh:bold', NS) is not None,
                    color=None if not color or color in ('#000000', 'none') else color)

    def parastyle(self, pid):
        p = self.paraprs.get(pid)
        st = dict(align='JUSTIFY', left=0, space_before=0, line=160, outline=False, tabs=[])
        if p is None:
            return st
        a = p.find('hh:align', NS)
        if a is not None:
            st['align'] = a.get('horizontal', 'JUSTIFY')
        h = p.find('hh:heading', NS)
        st['outline'] = h is not None and h.get('type') == 'OUTLINE'
        # HwpUnitChar 분기 안의 값이 실제 값이다
        case = p.find('.//hp:switch/hp:case', NS)
        node = case if case is not None else p
        m = node.find('.//hh:margin', NS)
        if m is not None:
            l, pv = m.find('hc:left', NS), m.find('hc:prev', NS)
            st['left'] = hu2pt(l.get('value', '0')) if l is not None else 0
            st['space_before'] = hu2pt(pv.get('value', '0')) if pv is not None else 0
        ls = node.find('.//hh:lineSpacing', NS)
        if ls is not None and ls.get('type') == 'PERCENT':
            st['line'] = int(ls.get('value', '160'))
        return st

    def borders(self, bfid):
        b = self.bfs.get(bfid)
        out, fill = dict(BORDER_NONE), None
        if b is None:
            return out, fill
        for side in ('left', 'right', 'top', 'bottom'):
            e = b.find('hh:%sBorder' % side, NS)
            if e is not None:
                w = float(re.sub(r'[^\d.]', '', e.get('width', '0.12')) or 0.12)
                out[side] = (e.get('type', 'NONE'), w, e.get('color', '#000000'))
        wb = b.find('.//hc:winBrush', NS)
        if wb is not None:
            fc = wb.get('faceColor', 'none')
            if fc and fc != 'none' and not fc.upper().startswith('#FF00000'):
                fill = fc if len(fc) == 7 else '#' + fc[-6:]
        return out, fill

    # --- 그림
    def image_path(self, ref):
        if self.tmpdir is None:
            self.tmpdir = os.path.join(os.environ.get('TEMP', '.'), 'hwpx_bindata')
            os.makedirs(self.tmpdir, exist_ok=True)
        for n in self.z.namelist():
            if n.startswith('BinData/') and os.path.splitext(os.path.basename(n))[0] == ref:
                out = os.path.join(self.tmpdir, os.path.basename(n))
                with open(out, 'wb') as f:
                    f.write(self.z.read(n))
                return out
        return None

    # --- 문단 → 블록들 (표·그림은 별도 블록으로 뒤에 붙는다)
    def para_blocks(self, p, in_cell=False):
        st = self.parastyle(p.get('paraPrIDRef'))
        pb = p.get('pageBreak') == '1'
        runs, extras = [], []
        for r in p.findall('hp:run', NS):
            cs = self.char(r.get('charPrIDRef'))
            for el in r:
                tag = el.tag.replace(HP, '')
                if tag == 't':
                    if el.text and not el.text.strip() and len(el.text) > 10:
                        runs.append(run('', tab=True, **cs))        # 공백으로 밀어 둔 오른쪽 정렬 → 탭
                    elif el.text:
                        runs.append(run(el.text, **cs))
                    for sub in el:
                        if sub.tag == HP + 'tab':
                            runs.append(run('', tab=True, **cs))
                            if sub.get('leader') not in (None, '0'):
                                st['leader'] = True             # 점선 채움 탭 (목차)
                        if sub.tail:
                            runs.append(run(sub.tail, **cs))
                elif tag == 'tbl':
                    extras.append(self.table_block(el, st['align']))
                elif tag == 'pic':
                    sz = el.find('hp:sz', NS); img = el.find('.//hc:img', NS)
                    path = self.image_path(img.get('binaryItemIDRef')) if img is not None else None
                    if path:
                        extras.append(pic(path, hu2pt(sz.get('width')), hu2pt(sz.get('height')), st['align']))
                elif tag == 'ctrl':
                    an = el.find('hp:autoNum', NS)
                    if an is not None:
                        runs.append(run('', field='PAGE' if an.get('numType') == 'PAGE' else 'NUMPAGES', **cs))
                elif tag == 'line' and not in_cell:
                    # 가로선 도형 → 문단 아래 테두리로 근사. 선 너비는 좌우 들여쓰기로 맞춘다
                    shape = el.find('hp:lineShape', NS); cur = el.find('hp:curSz', NS)
                    st['border_bottom'] = (shape.get('color', '#000000') if shape is not None else '#000000')
                    st['line_width'] = hu2pt(cur.get('width', '0')) if cur is not None else 0
        if st['outline'] and runs:
            self.outline += 1
            runs.insert(0, run(f'{self.outline}. ', **{k: runs[0][k] for k in ('font', 'size', 'bold', 'color')}))
        text_para = para(runs, st['align'], st['left'], st['space_before'], st['line'], pb, border_bottom=st.get('border_bottom'),
                         line_width=st.get('line_width'), leader=st.get('leader', False))
        segs = p.findall('hp:linesegarray/hp:lineseg', NS)
        if segs:   # 한글이 계산한 실제 줄 높이. Word 의 배수 줄간격은 글꼴에 따라 더 커져서 쪽이 밀린다
            text_para['line_pt'] = hu2pt(segs[0].get('vertsize', '0')) + hu2pt(segs[0].get('spacing', '0'))
        if extras and not runs and not st.get('border_bottom'):
            extras[0]['page_break'] = pb          # 표·그림만 든 문단은 빈 문단을 남기지 않는다
            return extras
        return [text_para] + extras

    def table_block(self, tbl, align):
        rows_n, cols_n = int(tbl.get('rowCnt')), int(tbl.get('colCnt'))
        grid = [[None] * cols_n for _ in range(rows_n)]
        widths = [0] * cols_n
        heights = [0] * rows_n
        for tr in tbl.findall('hp:tr', NS):
            for tc in tr.findall('hp:tc', NS):
                ad = tc.find('hp:cellAddr', NS); sp = tc.find('hp:cellSpan', NS); sz = tc.find('hp:cellSz', NS)
                c, r = int(ad.get('colAddr')), int(ad.get('rowAddr'))
                cspan, rspan = int(sp.get('colSpan', '1')), int(sp.get('rowSpan', '1'))
                bd, fill = self.borders(tc.get('borderFillIDRef'))
                sub = tc.find('hp:subList', NS)
                valign = sub.get('vertAlign', 'CENTER') if sub is not None else 'CENTER'
                blocks = []
                for pp in (sub.findall('hp:p', NS) if sub is not None else []):
                    blocks += self.para_blocks(pp, in_cell=True)
                grid[r][c] = cell(blocks, (cspan, rspan), bd, fill, valign, hu2pt(sz.get('width', '0')))
                if cspan == 1:
                    widths[c] = hu2pt(sz.get('width', '0'))
                if rspan == 1:
                    heights[r] = max(heights[r], hu2pt(sz.get('height', '0')))
        # 병합으로 너비를 못 받은 열은 나머지로 채운다
        total = hu2pt(tbl.find('hp:sz', NS).get('width')) if tbl.find('hp:sz', NS) is not None else sum(widths)
        missing = [i for i, w in enumerate(widths) if not w]
        if missing:
            rest = max(total - sum(widths), 0) / len(missing)
            for i in missing:
                widths[i] = rest
        t = table(grid, widths, align)
        t['row_heights'] = heights
        sz = tbl.find('hp:sz', NS)
        t['height'] = hu2pt(sz.get('height', '0')) if sz is not None else 0
        return t

    def parse(self):
        model = dict(page=self.page(), header=[], footer=[], blocks=[])
        for tag in ('header', 'footer'):
            for hf in self.sec.iter(HP + tag):
                for pp in hf.findall('.//hp:p', NS):
                    model[tag] += self.para_blocks(pp, in_cell=(tag == 'footer'))
                break
        # 본문 문단은 한글이 계산해 둔 줄 위치(lineseg, 쪽 위에서부터의 세로 위치)로 문단 앞 간격과
        # 쪽 넘김을 맞춘다 — 표지의 세로 배치와 "표지 다음 쪽부터 목차" 가 여기서 나온다
        prev_end = 0.0            # 직전 블록이 끝난 세로 위치 (pt). None 이면 모른다
        for p in self.sec.findall('hp:p', NS):
            blocks = self.para_blocks(p)
            segs = p.findall('hp:linesegarray/hp:lineseg', NS)
            if p.get('pageBreak') == '1':
                prev_end = 0.0
            if not segs or not blocks:
                model['blocks'] += blocks
                continue
            b0 = blocks[0]
            top = hu2pt(segs[0].get('vertpos', '0'))
            if prev_end is not None and top + 2 < prev_end and not b0.get('page_break'):
                b0['page_break'] = True      # 한글에서는 여기서 다음 쪽으로 넘어갔다
                prev_end = 0.0
            gap = top - prev_end if prev_end is not None else 0
            if b0['type'] in ('pic', 'table') and prev_end is not None:
                # 쪽 아래에 딱 붙는 그림·표(표지 로고)는 Word 에서 몇 pt 만 밀려도 다음 쪽으로 넘어간다 → 앞 간격을 줄여 여유를 둔다
                h0 = b0.get('h', 0) if b0['type'] == 'pic' else b0.get('height', 0)
                gap = min(gap, self.body_height() - 40 - prev_end - h0)
            if gap > 2:
                if b0['type'] == 'p':
                    if gap > b0['space_before'] + 2:
                        b0['space_before'] = gap
                elif b0['type'] == 'pic':
                    b0['space_before'] = gap
                elif b0['type'] == 'table' and gap > 14:
                    blocks.insert(0, para([], 'LEFT', space_before=gap - 12, line=100))
            if b0['type'] == 'p' and len(blocks) == 1:
                last = segs[-1]
                prev_end = hu2pt(last.get('vertpos', '0')) + hu2pt(last.get('vertsize', '0')) + hu2pt(last.get('spacing', '0'))
            else:
                h = sum(b.get('h', 0) if b['type'] == 'pic' else b.get('height', 0) if b['type'] == 'table' else 0 for b in blocks)
                prev_end = top + h if h else None
            model['blocks'] += blocks
        return model

    def body_height(self):
        pg = self.page()
        return pg['height'] - pg['top'] - pg['bottom']

    def page(self):
        pp = self.sec.find('.//hp:pagePr', NS); m = pp.find('hp:margin', NS)
        return dict(width=hu2pt(pp.get('width')), height=hu2pt(pp.get('height')),
                    left=hu2pt(m.get('left')), right=hu2pt(m.get('right')),
                    top=hu2pt(m.get('top')) + hu2pt(m.get('header')), bottom=hu2pt(m.get('bottom')) + hu2pt(m.get('footer')),
                    header=hu2pt(m.get('top')), footer=hu2pt(m.get('bottom')))


def parse_hwpx(path):
    return Hwpx(path).parse()


# ---------------------------------------------------------------- DOCX 렌더
def _set_font(r, font, size, bold, color):
    r.font.name = font
    rpr = r._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rf.set(qn(a), font)
    r.font.size = Pt(size)
    r.font.bold = bold
    if color:
        r.font.color.rgb = RGBColor.from_string(color.lstrip('#')[-6:].upper())


def _field(r, name):
    for t, txt in (('begin', None), (None, name), ('end', None)):
        if t:
            fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), t); r._element.append(fc)
        else:
            it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = f' {txt} '; r._element.append(it)


def _render_para(p, blk, body_width_pt):
    p.alignment = ALIGN.get(blk['align'], WD_ALIGN_PARAGRAPH.JUSTIFY)
    pf = p.paragraph_format
    pf.left_indent = Pt(blk['left']) if blk['left'] else None
    pf.space_before = Pt(blk['space_before']); pf.space_after = Pt(0)
    line_pt = blk.get('line_pt')
    if line_pt:      # 한글이 계산한 줄 높이 그대로 (표지처럼 세로 배치가 빡빡한 쪽이 밀리지 않게)
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY; pf.line_spacing = Pt(line_pt)
    else:
        sizes = [r['size'] for r in blk['runs'] if r.get('text') or r.get('field')]
        pf.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
        pf.line_spacing = Pt((max(sizes) if sizes else 10) * blk['line'] / 100.0)
    if blk.get('page_break'):
        pf.page_break_before = True
    ppr = p._element.get_or_add_pPr()
    for tag in ('w:autoSpaceDE', 'w:autoSpaceDN'):      # 한글과 숫자 사이에 Word 가 끼워 넣는 공백을 끈다 ("1단계" 유지)
        e = OxmlElement(tag); e.set(qn('w:val'), '0'); ppr.append(e)
    if any(r.get('tab') for r in blk['runs']):
        for tw in (4513, 9026):                          # 머리말 기본 탭(가운데·오른쪽)을 지운다
            pf.tab_stops.add_tab_stop(Emu(tw * 635), WD_TAB_ALIGNMENT.CLEAR)
        pf.tab_stops.add_tab_stop(Pt(body_width_pt), WD_TAB_ALIGNMENT.RIGHT,
                                  WD_TAB_LEADER.DOTS if blk.get('leader', True) else WD_TAB_LEADER.SPACES)
    if blk.get('border_bottom'):
        bd = OxmlElement('w:pBdr'); b = OxmlElement('w:bottom')
        b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '6'); b.set(qn('w:space'), '1')
        b.set(qn('w:color'), blk['border_bottom'].lstrip('#')[-6:]); bd.append(b); ppr.append(bd)
        lw = blk.get('line_width')
        if lw and lw < body_width_pt:
            pf.left_indent = pf.right_indent = Pt((body_width_pt - lw) / 2)
    for rd in blk['runs']:
        r = p.add_run()
        if rd.get('tab'):
            r.add_tab()
        elif rd.get('field'):
            _field(r, rd['field'])
        else:
            r.text = rd['text']
        _set_font(r, rd['font'], rd['size'], rd['bold'], rd.get('color'))


def _cell_borders(tc, borders, fill):
    tcpr = tc._element.get_or_add_tcPr()
    bd = OxmlElement('w:tcBorders')
    for side in ('top', 'left', 'bottom', 'right'):
        typ, w, color = borders.get(side, ('NONE', 0, '#000000'))
        e = OxmlElement('w:' + side)
        if typ == 'NONE':
            e.set(qn('w:val'), 'nil')
        else:
            e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(max(2, int(round(w / 0.125)))))   # 1/8pt 단위
            e.set(qn('w:color'), color.lstrip('#')[-6:])
        bd.append(e)
    tcpr.append(bd)
    if fill:
        shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), fill.lstrip('#')[-6:]); tcpr.append(shd)
    mar = OxmlElement('w:tcMar')
    for side, v in (('top', 20), ('bottom', 20), ('left', 60), ('right', 60)):
        e = OxmlElement('w:' + side); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); mar.append(e)
    tcpr.append(mar)


def _render_blocks(container, blocks, body_width_pt, doc):
    """container: Document / _Cell / header·footer. 셀 안에서는 첫 문단을 재활용한다."""
    first = container.paragraphs[0] if hasattr(container, 'paragraphs') and container.paragraphs and not container.paragraphs[0].text and len(container.paragraphs) == 1 and not isinstance(container, type(doc)) else None
    for blk in blocks:
        if blk['type'] == 'p':
            p = first if first is not None else container.add_paragraph()
            first = None
            _render_para(p, blk, body_width_pt)
        elif blk['type'] == 'pic':
            p = first if first is not None else container.add_paragraph()
            first = None
            p.alignment = ALIGN.get(blk['align'], WD_ALIGN_PARAGRAPH.CENTER)
            if blk.get('page_break'):
                p.paragraph_format.page_break_before = True
            if blk.get('space_before'):
                p.paragraph_format.space_before = Pt(blk['space_before'])
            p.add_run().add_picture(blk['path'], width=Pt(blk['w']), height=Pt(blk['h']))
        elif blk['type'] == 'table':
            if first is not None:
                first = None
            rows, cols = len(blk['rows']), len(blk['col_widths'])
            t = container.add_table(rows=rows, cols=cols)
            t.alignment = {'CENTER': WD_TABLE_ALIGNMENT.CENTER, 'LEFT': WD_TABLE_ALIGNMENT.LEFT}.get(blk['align'], WD_TABLE_ALIGNMENT.CENTER)
            t.autofit = False
            tblpr = t._element.tblPr
            lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblpr.append(lay)
            for i, w in enumerate(blk['col_widths']):
                for r in range(rows):
                    t.cell(r, i).width = Pt(w)
            for r, h in enumerate(blk.get('row_heights', [])):
                if 0 < h <= 30:                       # 한 줄짜리 행만. 상자 표의 높이는 내용이 정한다
                    t.rows[r].height = Pt(h)
            done = set()
            for r in range(rows):
                for c in range(cols):
                    cd = blk['rows'][r][c]
                    if cd is None or (r, c) in done:
                        continue
                    cs, rs = cd['span']
                    tc = t.cell(r, c)
                    if cs > 1 or rs > 1:
                        tc = tc.merge(t.cell(r + rs - 1, c + cs - 1))
                        for rr in range(r, r + rs):
                            for cc in range(c, c + cs):
                                done.add((rr, cc))
                    tc.vertical_alignment = {'TOP': WD_CELL_VERTICAL_ALIGNMENT.TOP, 'BOTTOM': WD_CELL_VERTICAL_ALIGNMENT.BOTTOM}.get(cd['valign'], WD_CELL_VERTICAL_ALIGNMENT.CENTER)
                    _cell_borders(tc, cd['borders'], cd['fill'])
                    _render_blocks(tc, cd['blocks'], body_width_pt, doc)
            if blk.get('page_break'):
                t.rows[0].cells[0].paragraphs[0].paragraph_format.page_break_before = True


def render(model, out_path):
    doc = Document()
    sec = doc.sections[0]
    pg = model['page']
    sec.page_width, sec.page_height = Pt(pg['width']), Pt(pg['height'])
    sec.left_margin, sec.right_margin = Pt(pg['left']), Pt(pg['right'])
    sec.top_margin, sec.bottom_margin = Pt(pg['top']), Pt(pg['bottom'])
    # Word 는 머리말·꼬리말 내용이 여백보다 크면 본문을 밀어낸다. 꼬리말(로고+한 줄 ≈ 40pt)이 들어갈 자리를 남긴다
    sec.header_distance = Pt(min(pg['header'], max(pg['top'] - 30, 5)))
    sec.footer_distance = Pt(min(pg['footer'], max(pg['bottom'] - 45, 5)))
    body_w = pg['width'] - pg['left'] - pg['right']
    normal = doc.styles['Normal']
    normal.font.name = '나눔고딕'; normal.font.size = Pt(10)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), '나눔고딕')
    normal.paragraph_format.space_after = Pt(0)
    if model.get('header'):
        sec.header.is_linked_to_previous = False
        _render_blocks(sec.header, model['header'], body_w, doc)
    if model.get('footer'):
        sec.footer.is_linked_to_previous = False
        _render_blocks(sec.footer, model['footer'], body_w, doc)
    _render_blocks(doc, model['blocks'], body_w, doc)
    doc.save(out_path)
    return out_path


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    src, dst = sys.argv[1], sys.argv[2]
    m = parse_hwpx(src)
    render(m, dst)
    print('saved', dst, 'blocks', len(m['blocks']))
