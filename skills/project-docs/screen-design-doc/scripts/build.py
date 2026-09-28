"""화면설계서 PPTX 와 클릭 데모 HTML 을 한 원본(spec.json + demo.html)에서 만든다.

python build.py 1.0.0
  → out/[1차] 화면 데모_v1.0.0.html
  → out/[1차] 화면설계서_v1.0.0.pptx
양식은 ../templates/ 에서 읽는다. 결과 폴더는 환경 변수 OUT_DIR 로 바꾼다(기본: 실행 폴더의 out).
spec.json 과 demo.html 은 예시 프로젝트(모임 관리 앱)의 내용이다. 자기 프로젝트에 맞게 바꿔 쓴다.
"""
import copy, json, os, re, sys
from pathlib import Path

from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Mm, Pt
from PIL import Image

HERE = Path(__file__).resolve().parent
TPL = HERE.parent / 'templates' / '[양식] V1.0_화면설계서.pptx'
OUTDIR = Path(os.environ.get('OUT_DIR') or Path.cwd() / 'out')
OUTDIR.mkdir(parents=True, exist_ok=True)
SHOTDIR = OUTDIR / '_build'          # 캡처 중간 파일
VER = sys.argv[1] if len(sys.argv) > 1 else '1.0.0'
FONT = '맑은 고딕'

SPEC = json.loads((HERE / 'spec.json').read_text(encoding='utf-8'))
ALL = SPEC['screens'] + SPEC['extra']
GROUPS = {g[0]: g for g in SPEC['groups']}

# 화면마다 UI 설계 장 구성. 장 하나 = 캡처 목록. 캡처 = (태그, 기기, 번호 제한, 설명)
# 기본은 스마트폰 앱과 웹을 나란히. PC 앱은 웹과 배치가 같아 폴더를 쓰는 E-02 에서만 따로 싣는다.
MOB, WEB = '스마트폰 앱', '웹 · PC 앱'
SHOTS = {
    'C-05': [[('C-05', 'mob', [1, 2, 3, 4, 5], MOB), ('C-05', 'web', [1, 2, 3, 4, 5], WEB)], [('C-05:empty', 'mob', [6], '후보가 없을 때')]],
    'C-06': [[('C-06', 'mob', [1, 2, 3, 4, 6], MOB + ' · 변경 뒤 다시 확정'), ('C-06', 'web', [1, 2, 3, 4, 6], WEB)], [('C-06:modal', 'mob', [5], '일정 변경 팝업')]],
    'E-01': [[('E-01', 'mob', [1, 2, 3, 4, 6], MOB), ('E-01', 'web', [1, 2, 3, 4, 6], WEB)], [('E-01:empty', 'mob', [5], '빈 분류')]],
    'E-02': [[('E-02', 'pc', [1, 2, 3, 4, 5], 'PC 앱 · 녹화 폴더')], [('E-02', 'mob', [1, 2, 3, 4, 5], MOB + ' · 파일 고르기'), ('E-02', 'web', [1, 2, 3, 4, 5], '웹 · Chrome·Edge 폴더 지정')],
             [('E-02:uploading', 'pc', [6, 7], 'PC 앱 · 올리는 중')]],
}
def plan(s):
    if s['id'][0] == 'R':
        return [[(s['id'], None, None, '')]]
    return SHOTS.get(s['id'], [[(s['id'], 'mob', None, MOB), (s['id'], 'web', None, WEB)]])


# ---------------------------------------------------------------- 1. 데모 HTML
def build_demo():
    html = (HERE / 'demo.html').read_text(encoding='utf-8')
    html = html.replace('/*SPEC*/null', json.dumps(SPEC, ensure_ascii=False))
    html = html.replace('화면 데모 v1', f'화면 데모 v{VER}')
    out = OUTDIR / f'[1차] 화면 데모_v{VER}.html'
    out.write_text(html, encoding='utf-8')
    css = re.search(r'<style>(.*?)</style>', html, re.S).group(1)
    extra = (HERE / 'extra.html').read_text(encoding='utf-8').replace('/*CSS*/', css)
    SHOTDIR.mkdir(exist_ok=True)
    (SHOTDIR / 'extra.html').write_text(extra, encoding='utf-8')
    (SHOTDIR / 'demo.html').write_text(html, encoding='utf-8')
    return out


# ---------------------------------------------------------------- 2. 캡처
def capture():
    shots = {}
    with sync_playwright() as p:
        b = p.chromium.launch()
        pg = b.new_page(viewport={'width': 1500, 'height': 1000}, device_scale_factor=2)
        errs = []
        pg.on('pageerror', lambda e: errs.append(str(e)))
        demo, extra = (SHOTDIR / 'demo.html').as_uri(), (SHOTDIR / 'extra.html').as_uri()
        for s in ALL:
            shots[s['id']] = []
            for slide in plan(s):
                part = []
                for tag, dev, only, label in slide:
                    q = f'?shot&wire&marks&t={tag}' + (f'&dev={dev}' if dev else '') + (f'&only={",".join(map(str, only))}' if only else '')
                    pg.goto((extra if s['id'][0] == 'R' else demo) + q + '#' + tag)
                    pg.evaluate('document.fonts.ready')
                    pg.wait_for_timeout(500)
                    f = SHOTDIR / f'{tag.replace(":", "_")}_{dev or "x"}.png'
                    pg.locator('#device').screenshot(path=str(f))
                    marks = pg.eval_on_selector_all('#device .mk', 'e => e.map(x => +x.textContent)')
                    part.append((f, label, sorted(marks)))
                shots[s['id']].append(part)
        b.close()
    if errs:
        raise SystemExit('데모 스크립트 오류: ' + ' / '.join(errs[:3]))
    return shots


# ---------------------------------------------------------------- 3. PPTX 도우미
def set_text(tf_or_cell, lines, size=None, bold=None, color=None):
    """첫 문단과 첫 글자 서식을 살려 두고 글만 바꾼다. lines 는 문자열 또는 목록."""
    tf = tf_or_cell.text_frame
    if isinstance(lines, str):
        lines = [lines]
    paras = tf.paragraphs
    p0 = paras[0]._p
    ppr = copy.deepcopy(p0.find(qn('a:pPr')))
    r0 = p0.find(qn('a:r'))
    rpr = copy.deepcopy(r0.find(qn('a:rPr'))) if r0 is not None and r0.find(qn('a:rPr')) is not None else None
    for p in paras[1:]:
        p._p.getparent().remove(p._p)
    for el in list(p0):
        if el.tag != qn('a:pPr'):
            p0.remove(el)
    for i, line in enumerate(lines):
        if i == 0:
            p = p0
        else:
            p = copy.deepcopy(p0)
            for el in list(p):
                if el.tag != qn('a:pPr'):
                    p.remove(el)
            p0.getparent().append(p)
        r = p.makeelement(qn('a:r'), {})
        if rpr is not None:
            r.append(copy.deepcopy(rpr))
        t = r.makeelement(qn('a:t'), {})
        t.text = line
        r.append(t)
        p.append(r)
    for p in tf.paragraphs:
        for r in p.runs:
            if size:
                r.font.size = Pt(size)
            if bold is not None:
                r.font.bold = bold
            if color:
                r.font.color.rgb = RGBColor.from_string(color)
            r.font.name = FONT
            rp = r._r.get_or_add_rPr()
            for tag in ('a:ea',):
                e = rp.find(qn(tag))
                if e is None:
                    e = rp.makeelement(qn(tag), {})
                    rp.append(e)
                e.set('typeface', FONT)


def drop_slide(prs, idx):
    sld = prs.slides._sldIdLst[idx]
    prs.part.drop_rel(sld.get(qn('r:id')))
    prs.slides._sldIdLst.remove(sld)


def layout(prs, name):
    return next(l for l in prs.slide_layouts if l.name == name)


def page_no(slide, src_box, n):
    el = copy.deepcopy(src_box._element)
    slide.shapes._spTree.append(el)
    box = slide.shapes[-1]
    set_text(box, str(n))


def clone_table(slide, src_frame, ncols, nrows, left, top, widths_mm, row_h_mm=None):
    """양식의 표를 복제해 열과 행 수를 맞춘다. 1행은 머리, 2행을 본문 행 틀로 쓴다."""
    el = copy.deepcopy(src_frame._element)
    slide.shapes._spTree.append(el)
    fr = slide.shapes[-1]
    tbl = fr._element.graphic.graphicData.tbl
    grid = tbl.tblGrid
    cols = grid.findall(qn('a:gridCol'))
    trs = tbl.findall(qn('a:tr'))
    # 열 맞추기
    while len(cols) < ncols:
        grid.append(copy.deepcopy(cols[-1]))
        for tr in trs:
            tcs = tr.findall(qn('a:tc'))
            tcs[-1].addnext(copy.deepcopy(tcs[-1]))
        cols = grid.findall(qn('a:gridCol'))
    while len(cols) > ncols:
        grid.remove(cols[-1])
        for tr in trs:
            tr.remove(tr.findall(qn('a:tc'))[-1])
        cols = grid.findall(qn('a:gridCol'))
    # 행 맞추기
    body = trs[1]
    for tr in trs[1:]:
        tbl.remove(tr)
    for _ in range(nrows - 1):
        tbl.append(copy.deepcopy(body))
    for gc, w in zip(grid.findall(qn('a:gridCol')), widths_mm):
        gc.set('w', str(int(Mm(w))))
    if row_h_mm:
        for tr in tbl.findall(qn('a:tr')):
            tr.set('h', str(int(Mm(row_h_mm))))
    fr.left, fr.top = Mm(left), Mm(top)
    fr.width = Mm(sum(widths_mm))
    return fr.table


def fill_table(table, rows, size=8.5, head=9, left=()):
    """left 에 든 열 번호는 왼쪽 정렬, 나머지는 가운데. 양식 표에서 딸려 온 글머리 기호는 뺀다."""
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            cell = table.cell(i, j)
            set_text(cell, v, size=head if i == 0 else size, bold=True if i == 0 else None)
            for p in cell.text_frame.paragraphs:
                ppr = p._p.get_or_add_pPr()
                for tag in ('a:buChar', 'a:buAutoNum', 'a:buFont', 'a:buNone'):
                    for e in ppr.findall(qn(tag)):
                        ppr.remove(e)
                ppr.append(ppr.makeelement(qn('a:buNone'), {}))
                ppr.set('marL', '0'); ppr.set('indent', '0')
                p.alignment = PP_ALIGN.LEFT if (j in left and i > 0) else PP_ALIGN.CENTER


def textbox(slide, x, y, w, h, text, size=9, bold=False, color='333333', align=None):
    tb = slide.shapes.add_textbox(Mm(x), Mm(y), Mm(w), Mm(h))
    tf = tb.text_frame
    tf.word_wrap = True
    lines = text if isinstance(text, list) else text.split('\n')
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = line
        r.font.size, r.font.bold, r.font.name = Pt(size), bold, FONT
        r.font.color.rgb = RGBColor.from_string(color)
        r._r.get_or_add_rPr().append(r._r.makeelement(qn('a:ea'), {'typeface': FONT}))
        if align:
            p.alignment = align
    return tb


# ---------------------------------------------------------------- 4. 흐름도 도형
KIND = {'T': MSO_SHAPE.FLOWCHART_TERMINATOR, 'P': MSO_SHAPE.FLOWCHART_PROCESS, 'D': MSO_SHAPE.FLOWCHART_DECISION,
        'C': MSO_SHAPE.FLOWCHART_CONNECTOR, 'DOC': MSO_SHAPE.FLOWCHART_DOCUMENT, 'IN': MSO_SHAPE.FLOWCHART_DATA,
        'DB': MSO_SHAPE.FLOWCHART_MAGNETIC_DISK}


def node(slide, kind, x, y, w, h, text, size=8):
    sh = slide.shapes.add_shape(KIND[kind], Mm(x), Mm(y), Mm(w), Mm(h))
    sh.shadow.inherit = False
    fill = {'T': 'D9D9D9', 'D': 'EDEDED', 'C': '262626', 'DB': 'F2F2F2'}.get(kind, 'FFFFFF')
    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    sh.line.color.rgb = RGBColor(0x40, 0x40, 0x40); sh.line.width = Pt(0.9)
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = Mm(0.8); tf.margin_top = tf.margin_bottom = Mm(0.3)
    tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.name = FONT
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if kind == 'C' else RGBColor(0x1A, 0x1A, 0x1A)
    r._r.get_or_add_rPr().append(r._r.makeelement(qn('a:ea'), {'typeface': FONT}))
    return sh


def arrow(slide, x1, y1, x2, y2, elbow=False, color='404040'):
    c = slide.shapes.add_connector(MSO_CONNECTOR.ELBOW if elbow else MSO_CONNECTOR.STRAIGHT, Mm(x1), Mm(y1), Mm(x2), Mm(y2))
    c.line.color.rgb = RGBColor.from_string(color); c.line.width = Pt(0.9)
    ln = c.line._get_or_add_ln()
    ln.append(ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'sm', 'len': 'sm'}))
    return c


def lane_flow(slide, title, steps, x0=12, y0=38, total_w=315, label_w=32, h=128, nw=34, nh=8, gap=5, notes=None):
    """STEP 열 흐름도. 양식 User flow 장의 표 모양을 그대로 따른다."""
    cw = (total_w - label_w) / len(steps)
    # 틀
    cols = [(x0 + label_w + i * cw, cw, s[0]) for i, s in enumerate(steps)]
    if title:
        cols = [(x0, label_w, '')] + cols
    for (x, w, head) in cols:
        hd = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Mm(x), Mm(y0), Mm(w), Mm(9))
        hd.fill.solid(); hd.fill.fore_color.rgb = RGBColor(0xD9, 0xD9, 0xD9); hd.line.color.rgb = RGBColor(0xBF, 0xBF, 0xBF)
        hd.shadow.inherit = False
        bd = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Mm(x), Mm(y0 + 9), Mm(w), Mm(h))
        bd.fill.background(); bd.line.color.rgb = RGBColor(0xBF, 0xBF, 0xBF); bd.shadow.inherit = False
        if head:
            textbox(slide, x + 1.5, y0 + 1.6, w - 3, 6, head, size=9, bold=True, color='1A1A1A')
    if title:
        textbox(slide, x0 + 1, y0 + 9 + h / 2 - 6, label_w - 2, 12, title, size=9.5, bold=True, color='1A1A1A', align=PP_ALIGN.CENTER)
    ends = []
    for i, (_, items) in enumerate(steps):
        cx = x0 + label_w + i * cw + cw / 2 - (4 if any(it.get('note') for it in items) else 0)
        y = y0 + 9 + 5
        pos = []
        for it in items:
            k = it['k']
            w_, h_ = (8, 8) if k == 'C' else (nw + 4, nh + 3) if k == 'D' else (nw, nh)
            node(slide, k, cx - w_ / 2, y, w_, h_, it['t'], size=7 if k == 'C' else 7.5)
            pos.append((cx, y, w_, h_))
            if it.get('note'):
                textbox(slide, cx + w_ / 2 + 1, y - 0.5, cw / 2 - w_ / 2 + 2, h_ + 2, it['note'], size=6.5, color='595959')
            y += h_ + gap
        for j, it in enumerate(items):
            cx_, yy, w_, h_ = pos[j]
            if j + 1 < len(items) and it.get('arrow', True):
                arrow(slide, cx_, yy + h_, cx_, pos[j + 1][1])
                if it['k'] == 'D':
                    textbox(slide, cx_ + 0.5, yy + h_ - 0.8, 10, 4, 'YES', size=6.5, bold=True)
            if it['k'] == 'D' and it.get('no') is not None:
                tx, ty, tw, th = pos[it['no']]
                xr = cx_ + w_ / 2
                side = xr + 3
                ln1 = arrow(slide, xr, yy + h_ / 2, side, yy + h_ / 2); ln1.line._get_or_add_ln().remove(ln1.line._get_or_add_ln().find(qn('a:tailEnd')))
                ln2 = arrow(slide, side, yy + h_ / 2, side, ty + th / 2); ln2.line._get_or_add_ln().remove(ln2.line._get_or_add_ln().find(qn('a:tailEnd')))
                arrow(slide, side, ty + th / 2, tx + tw / 2, ty + th / 2)
                textbox(slide, xr - 0.5, yy + h_ / 2 - 4.5, 10, 4, 'NO', size=6.5, bold=True)
        ends.append((pos[0], pos[-1] if items[-1].get('out', True) else pos[-2]))
    # 열 사이 연결
    for i in range(len(steps) - 1):
        (_, last) = ends[i][0], ends[i][1]
        first = ends[i + 1][0]
        lx, ly, lw, lh = last
        fx, fy, fw, fh = first
        mid = x0 + label_w + (i + 1) * cw - 2
        a = arrow(slide, lx + lw / 2, ly + lh / 2, mid, ly + lh / 2); a.line._get_or_add_ln().remove(a.line._get_or_add_ln().find(qn('a:tailEnd')))
        b = arrow(slide, mid, ly + lh / 2, mid, fy + fh / 2); b.line._get_or_add_ln().remove(b.line._get_or_add_ln().find(qn('a:tailEnd')))
        arrow(slide, mid, fy + fh / 2, fx - fw / 2, fy + fh / 2)


def N(k, t, **kw):
    return dict(k=k, t=t, **kw)


FLOWS = [
    ('일정 조율\nUser flow', [
        ('1 STEP. 요청', [N('T', 'START'), N('P', '팀장이 일정 요청 만들기'), N('P', '대상 주와 진행 예상 시간 정하기'), N('P', '요청 보내기'), N('P', '참여자에게 알림 (앱, 디스코드)')]),
        ('2 STEP. 가능 시간 입력', [N('P', '안 되는 날, 바뀔 수도 있는 날 표시'), N('P', '되는 날의 시간 정하기'), N('P', '제출'), N('D', '전원 제출?', no=5), N('P', '후보 시간 산출'), N('P', '마감 전 재촉 알림', arrow=False, out=False)]),
        ('3 STEP. 투표', [N('P', '후보 시간 알림'), N('P', '후보에 투표 (여러 개)'), N('D', '전원 투표?', no=4), N('P', '최다 득표로 확정'), N('P', '투표 기다리기', arrow=False, out=False)]),
        ('4 STEP. 확정과 변경', [N('P', '디스코드 자동 공지'), N('P', '확정 일정 보기'), N('D', '팀장이 변경 시작?', no=6), N('P', '이전 확정 취소'), N('P', '전원 알림'), N('C', '3', arrow=False), N('T', 'END', out=True)]),
    ]),
    ('공지\nUser flow', [
        ('1 STEP. 작성', [N('T', 'START'), N('P', '공지 쓰기'), N('P', '템플릿 고르기'), N('P', '빈칸 채우기'), N('P', '대상 고르기')]),
        ('2 STEP. 저장 또는 발송', [N('P', '미리보기'), N('D', '바로 보낼까?', no=4), N('P', '발송하기, 확인 팝업'), N('P', '공지방으로 전송', arrow=False), N('P', '임시 저장하기', out=False)]),
        ('3 STEP. 결과', [N('D', '전송 성공?', no=3), N('P', '목록에 발송 표시'), N('T', 'END', arrow=False), N('P', '목록에 실패 표시'), N('P', '다시 보내기', out=False)]),
    ]),
    ('자료 올리기\nUser flow', [
        ('1 STEP. 준비 (어느 기기든)', [N('T', 'START'), N('D', '폴더 지정됨?', no=4), N('P', '지정 폴더 목록 보기'), N('P', '새로 고침', arrow=False), N('P', '폴더 지정 또는 파일 고르기 (R-05)', arrow=False)]),
        ('2 STEP. 올리기', [N('P', '파일 고르기 또는 끌어다 놓기'), N('P', '확인 팝업, 분류 고르기'), N('P', '올리기와 진행률'), N('D', '끝까지 올라감?', no=5), N('P', '자료실에 표시'), N('P', '이어 올리기', arrow=False, out=False)]),
        ('3 STEP. 공유', [N('P', '자료 올림 공지 (선택)'), N('P', '회원이 자료실에서 보기'), N('P', '내려받기'), N('T', 'END')]),
    ]),
    ('시작과 로그인\nUser flow', [
        ('1 STEP. 앱 열기', [N('T', 'START'), N('D', '로그인 상태?', no=3), N('P', '홈 (B-01)', arrow=False), N('P', '로그인 (A-01)')]),
        ('2 STEP. 로그인', [N('D', '계정 있음?', no=4), N('P', '이메일, 비밀번호'), N('D', '맞음?', no=1, note='· 비밀번호를 잊으면\n  비밀번호 찾기 (A-03)'), N('P', '홈으로', arrow=False), N('P', '회원가입, 인증 메일 (A-02)', out=False)]),
        ('3 STEP. 쓰기', [N('P', '홈에서 할 일 보기'), N('P', '하단 탭으로 이동'), N('P', '내 정보, 로그아웃'), N('T', 'END')]),
    ]),
]

LOGICS = [
    ('Logic process · 후보 시간 산출과 확정', [
        ('입력', [N('T', 'Start'), N('DOC', '가능 시간 입력 (C-03, C-04)'), N('IN', '날 상태, 시간대', note='· 날마다 되는 날, 안 됨, 바뀔 수도\n· 시작, 끝 (30분 단위)'), N('D', '되는 날 있음?', no=1), N('DB', '가능시간 DB')]),
        ('산출', [N('D', '전원 제출?', no=4), N('P', '날마다 전원 시간대의 교집합'), N('D', '교집합 ≥ 진행 예상 시간?', no=5, note='· 진행 예상 시간은 팀장이 정함\n· 모자라면 후보에서 뺀다'), N('P', '후보로 저장', note='· 바뀔 수도 있는 날이 끼면 2순위\n· 1순위 먼저, 같으면 날짜순'), N('P', '마감까지 대기', arrow=False, out=False), N('P', '그 날은 뺀다', arrow=False, out=False)]),
        ('투표와 확정', [N('DB', '후보시간 DB'), N('DOC', '투표 (C-05)'), N('D', '전원 투표?', no=5), N('P', '최다 득표 확정', note='· 동점이면 날짜가 빠른 쪽\n· 후보 0개면 다시 요청'), N('DB', '확정일정 DB', arrow=False), N('P', '마감 뒤 확정 (미정)', arrow=False, out=False)]),
        ('알림과 변경', [N('P', '디스코드 자동 공지'), N('D', '변경 시작?', no=4), N('P', '확정을 변경으로 취소', note='· 사유를 남긴다\n· 전원에게 알림'), N('P', '회차 올리고 다시 투표', arrow=False), N('T', 'End')]),
    ]),
    ('Logic process · 공지 발송', [
        ('작성', [N('T', 'Start'), N('DOC', '공지 작성 (D-02)'), N('IN', '템플릿, 빈칸, 대상', note='· 빈칸은 {이름}\n· 대상은 역할로 제한'), N('D', '역할 확인 (서버)', no=4), N('P', '권한 없음 (R-04)', arrow=False, out=False)]),
        ('저장', [N('D', '발송하기?', no=3), N('P', '상태 대기로 저장'), N('P', '대상에 맞는 공지방 찾기', arrow=False, note='· 전체, 부팀장·팀원: 팀 공지방\n· 팀장, 팀장·부팀장: 운영진\n· 특정인: 미정'), N('P', '임시저장으로 저장', out=False), N('DB', '공지 DB', arrow=False)]),
        ('전송', [N('P', '디스코드로 전송 (웹훅)'), N('D', '10초 안 성공?', no=3), N('P', '상태 발송', arrow=False), N('P', '상태 실패, 목록에 표시'), N('P', '다시 보내기', out=False)]),
    ]),
]


# ---------------------------------------------------------------- 5. PPTX 만들기
def build_pptx(shots):
    prs = Presentation(str(TPL))
    s_title, s_hist, s_over, _, _, _, s_desc = list(prs.slides)
    pno_src = next(sh for sh in s_hist.shapes if sh.shape_type == 17)
    hist_tbl = next(sh for sh in s_hist.shapes if sh.has_table)
    desc_tbl = next(sh for sh in s_desc.shapes if sh.has_table)
    desc_pno = next(sh for sh in s_desc.shapes if sh.shape_type == 17)

    # 표지
    ph = {p.placeholder_format.idx: p for p in s_title.placeholders}
    set_text(ph[0], SPEC['title'])
    vals = [f'v{VER}', SPEC['history'][-1][1].replace('/', '.'), SPEC['writer']]
    for p, v in zip(sorted([p for p in s_title.placeholders if p.placeholder_format.idx != 0], key=lambda x: x.left), vals):
        set_text(p, v)

    # History
    t = hist_tbl.table
    for i, row in enumerate(SPEC['history'], start=1):
        for j, v in enumerate(row):
            set_text(t.cell(i, j), v)

    # 서비스 개요
    ot = next(sh for sh in s_over.shapes if sh.has_table)
    tbl = ot.table._tbl
    trs = tbl.findall(qn('a:tr'))
    for i, (k, lines) in enumerate(SPEC['overview'], start=1):
        set_text(ot.table.cell(i, 1), lines)
    for tr in trs[6:]:
        tbl.remove(tr)
    textbox(s_over, 33, 150, 272, 10,['기획일정은 WBS 일정표를 따른다.',
                                        '1차 기간 2026-09-21 ~ 2026-11-20. 와이어프레임 초안 홍길동, 화면 검수 이영희.'], size=9, color='404040')

    n = 4
    def new(name):
        nonlocal n
        sl = prs.slides.add_slide(layout(prs, name))
        if name != '사용자 지정 레이아웃':
            page_no(sl, pno_src if name == 'Summary' else desc_pno, n)
        n += 1
        return sl

    def summary(title):
        sl = new('Summary')
        set_text(sl.shapes.title, title)
        return sl

    # 화면 목록
    rows = [['화면 ID', '화면명', '그룹', '플랫폼', '관련 요구사항', '데모']]
    for s in ALL:
        rows.append([s['id'], s['title'], GROUPS[s['group']][1], s['platform'], s['fr'], '없음' if s['id'][0] == 'R' else '있음'])
    half = (len(rows) - 1 + 1) // 2
    for k, part in enumerate([rows[1:1 + half], rows[1 + half:]]):
        sl = summary(f'화면 목록 ({k + 1}/2)')
        tb = clone_table(sl, hist_tbl, 6, len(part) + 1, 20, 38, [22, 70, 26, 30, 110, 16], row_h_mm=8.6)
        fill_table(tb, [rows[0]] + part, size=8.5, left=(1, 4))
    # 공통 가이드
    sl = summary('공통 가이드 · 내비게이션과 역할별 분기')
    textbox(sl, 20, 34, 150, 6, '내비게이션', size=10, bold=True, color='1A1A1A')
    nav = SPEC['common']['nav']
    tb = clone_table(sl, hist_tbl, 2, len(nav) + 1, 20, 41, [34, 262], row_h_mm=7.6)
    fill_table(tb, [['구분', '규칙']] + nav, left=(1,))
    textbox(sl, 20, 86, 150, 6, '역할별로 보이는 기능 (분기처리)', size=10, bold=True, color='1A1A1A')
    r = SPEC['common']['roles']
    tb = clone_table(sl, hist_tbl, 5, len(r), 20, 93, [96, 50, 50, 50, 50], row_h_mm=7)
    fill_table(tb, r, left=(0,))
    textbox(sl, 20, 153, 296, 10, SPEC['common']['roles_note'], size=9, color='404040')
    sl = summary('공통 가이드 · 알림 정책과 공통 상태')
    textbox(sl, 20, 34, 150, 6, '알림 정책', size=10, bold=True, color='1A1A1A')
    r = SPEC['common']['notify']
    tb = clone_table(sl, hist_tbl, 4, len(r), 20, 41, [80, 70, 40, 106], row_h_mm=8)
    fill_table(tb, r, left=(0, 1, 3))
    textbox(sl, 20, 98, 296, 40, [
        '알림은 앱 안의 알림 목록과 디스코드 공지방 두 곳이다. 휴대폰과 PC 알림은 디스코드 앱이 띄운다(R-02).',
        '알림을 한 기기에서 읽으면 다른 기기에서도 읽음이 된다.',
        '팝업은 확인형, 위험 확인형, 입력형, 선택형, 토스트 다섯 가지만 쓴다(R-03).',
        '목록 화면은 불러오는 중, 빈 목록, 오류, 권한 없음 네 상태를 가진다(R-04).',
        '되돌리기 어려운 동작(일정 변경, 로그아웃)은 빨간 글자 버튼과 위험 확인형 팝업을 쓴다.',
    ], size=9.5, color='333333')

    # User flow, Logic process
    for title, steps in FLOWS:
        sl = summary('User flow')
        lane_flow(sl, title, steps)
    for title, steps in LOGICS:
        sl = summary(title)
        lane_flow(sl, '', steps, label_w=0, nw=36)

    # 메뉴 그룹 + UI 설계
    for g in SPEC['groups']:
        gid, gname, gdesc = g
        sl = new('사용자 지정 레이아웃')
        ph = {p.placeholder_format.idx: p for p in sl.placeholders}
        set_text(ph[0], f'{gid}. {gname}')
        set_text(ph[10], gdesc)
        for s in [x for x in ALL if x['group'] == gid]:
            parts = shots[s['id']]
            for pi, part in enumerate(parts):
                desc_slide(new('Description'), desc_tbl, s, gid, gname, part, f' ({pi + 1}/{len(parts)})' if len(parts) > 1 else '')

    # 양식 예시 장표(User flow, Logic process, 메뉴 그룹, Description 예시)를 뺀다
    for idx in (6, 5, 4, 3):
        drop_slide(prs, idx)
    # 복제한 도형의 id 가 겹치면 PowerPoint 가 파일을 열지 못한다. 슬라이드마다 새로 매긴다.
    for sl in prs.slides:
        for k, el in enumerate(sl.shapes._spTree.iter(qn('p:cNvPr')), start=2):
            el.set('id', str(k))
    out = OUTDIR / f'[1차] 화면설계서_v{VER}.pptx'
    prs.save(str(out))
    return out


def desc_slide(sl, desc_tbl, s, gid, gname, shot_items, suffix):
    """UI 설계 장 하나. 위 제목 두 칸, 오른쪽 Description 표, 왼쪽 캡처."""
    phs = sorted(sl.placeholders, key=lambda p: p.left)
    set_text(phs[0], f"{s['id']} {s['title']}{suffix}")
    set_text(phs[1], f"{gid}. {gname} · {s['platform']}")
    sl.shapes._spTree.append(copy.deepcopy(desc_tbl._element))
    dt = sl.shapes[-1].table
    set_text(dt.cell(1, 0), [s['summary'], f"관련 요구사항 {s['fr']}"], size=8)
    items = s['items']
    trs = dt._tbl.findall(qn('a:tr'))
    for extra_tr in trs[2 + len(items):]:
        dt._tbl.remove(extra_tr)
    for i, it in enumerate(items):
        set_text(dt.cell(2 + i, 1), it, size=8)
        trs[2 + i].set('h', str(int(Mm(11 if len(it) > 60 else 8.5))))
    place_shots(sl, shot_items)


def place_shots(sl, items):
    """왼쪽 화면 자리(약 232 × 158 mm)에 캡처를 나란히 넣는다."""
    X0, Y0, W, H = 6, 17, 230, 158
    sizes = [Image.open(f).size for f, _, _ in items]
    cap = 6 if any(label for _, label, _ in items) else 0
    gap = 6
    k = len(items)
    # 나란히 놓기와 위아래로 쌓기 중 캡처가 더 크게 들어가는 쪽을 고른다
    side = min((H - cap) / max(h for _, h in sizes), (W - gap * (k - 1)) / sum(w for w, _ in sizes))
    stack = min(W / max(w for w, _ in sizes), (H - (cap + gap) * k + gap) / sum(h for _, h in sizes))
    if k > 1 and stack > side:
        y = Y0
        for (f, label, _), (w, h) in zip(items, sizes):
            x = X0 + (W - w * stack) / 2
            if label:
                textbox(sl, x, y, w * stack, 5, label, size=8.5, bold=True, color='404040')
            pic = sl.shapes.add_picture(str(f), Mm(x), Mm(y + cap), Mm(w * stack), Mm(h * stack))
            pic.line.color.rgb = RGBColor(0xD0, 0xD0, 0xD0); pic.line.width = Pt(0.5)
            y += cap + h * stack + gap
        return
    total = sum(w for w, _ in sizes) * side + gap * (k - 1)
    x = X0 + (W - total) / 2
    for (f, label, _), (w, h) in zip(items, sizes):
        pic = sl.shapes.add_picture(str(f), Mm(x), Mm(Y0 + cap), Mm(w * side), Mm(h * side))
        pic.line.color.rgb = RGBColor(0xD0, 0xD0, 0xD0); pic.line.width = Pt(0.5)
        if label:
            textbox(sl, x, Y0, w * side, 5, label, size=8.5, bold=True, color='404040')
        x += w * side + gap


# ---------------------------------------------------------------- 6. 점검
def check(shots):
    # 필수 FR 목록은 요구사항서 생성기(planning-docs 스킬)의 BODY 에서 읽는다. REQ_BUILD 로 다른 파일을 가리킬 수 있다.
    rb = Path(os.environ.get('REQ_BUILD') or HERE.parents[1] / 'planning-docs' / 'scripts' / 'req_build.py')
    if not rb.exists():
        print('요구사항서 생성기를 찾지 못해 FR 추적 점검을 건너뛴다:', rb)
        must = []
    else:
        must = re.findall(r"\['(FR-\d+)', '[^']*', '필수'\]", rb.read_text(encoding='utf-8'))
    cover = set()
    for s in ALL:
        for a, b in re.findall(r'FR-(\d+)(?:\s*~\s*FR-(\d+))?', s['fr']):
            cover.update(range(int(a), int(b or a) + 1))
    missing = [f for f in must if int(f[3:]) not in cover]
    print('필수 FR', len(must), '개 중 화면에 걸린 것', len(must) - len(missing), '/ 빠진 것', missing or '없음')
    for s in ALL:
        got = sorted({m for part in shots[s['id']] for _, _, ms in part for m in ms})
        want = list(range(1, len(s['items']) + 1))
        if got != want:
            print('  번호 불일치', s['id'], '캡처', got, '설명', want)


if __name__ == '__main__':
    demo = build_demo()
    shots = capture()
    check(shots)
    pptx = build_pptx(shots)
    print('저장', demo.name, '/', pptx.name)
