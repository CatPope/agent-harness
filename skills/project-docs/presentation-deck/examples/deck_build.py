# 참고용 예시다. 이 파일은 한 프로젝트(모임 관리 앱)의 발표 생성기를 가명으로 옮긴 것이라 그대로는 돌지 않는다.
# 화면 캡처(화면설계서 데모), 구성 그림, 썸네일 폴더를 전제한다. set_text · duplicate · move · picture_fit 같은
# 양식 채우기 함수와 장 구성 방식을 가져다 쓴다. 양식 원본은 ../templates/ 에 있다.
"""프로젝트 제안 발표 PPT 생성기. 양식 docs/양식/[양식] 프로젝트 제안서.pptx 를 채운다.

python .claude/tools/docgen/deck_build.py 1.0.0
  → docs/발표/[1차] 프로젝트 제안 발표_v<버전>.pptx
스킬: .claude/skills/presentation-deck

- 양식 슬라이드는 도형 id 로 글을 바꾼다. 양식이 바뀌면 id 부터 다시 확인한다(스킬 참고).
- 그림(흐름·비교·단계 막대·이니셜)은 matplotlib 으로 양식 색에 맞춰 그린다.
- 화면 캡처는 화면 데모 HTML(.claude/tools/docgen/screen/_build/demo.html)에서 컬러로 찍는다.
- 새 슬라이드(화면 미리보기, 시스템 구성)는 제안 솔루션 장을 복제해 틀만 남긴다.
"""
import copy, os, sys, re
from pathlib import Path

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, Circle
from PIL import Image
from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Mm, Pt, Emu
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

sys.stdout.reconfigure(encoding='utf-8')
ROOT = Path(__file__).resolve().parents[3]
TPL = ROOT / 'docs' / '양식' / '[양식] 프로젝트 제안서.pptx'
VER = sys.argv[1] if len(sys.argv) > 1 else '1.0.0'
OUT = ROOT / 'docs' / '발표' / f'[1차] 프로젝트 제안 발표_v{VER}.pptx'
WORK = Path(__file__).resolve().parent / '_deck'
DEMO = Path(__file__).resolve().parent / 'screen' / '_build' / 'demo.html'
ARCH = ROOT / 'docs' / '설계' / '[1차] 구성 아키텍처_v1.2.0.png'
REPO = 'github.com/example-app'
REPO_SHORT = 'example-app'   # 꼬리말 알약은 폭이 좁다
# 이 양식은 508 x 286 mm 로 보통 슬라이드의 1.5 배다. 글자 크기도 1.5 배로 잡는다(본문 22~26pt).

DARK, OLIVE, GRAY, LIGHT, INK = '#575849', '#828A73', '#97988E', '#DEDEDC', '#3A3A3A'
plt.rcParams['font.family'] = 'Malgun Gothic'
plt.rcParams['axes.unicode_minus'] = False

# ---------------------------------------------------------------- 내용 (짧은 단문만)
CONTENTS = ['프로젝트 소개', '문제 인식', '현황 분석', '제안 솔루션', '프로젝트 계획', '기대 효과', '위험 관리', '팀 소개', '현황']
S3 = {33: '무엇을 만드나', 32: '무엇을 이루나',
      # 줄 사이 빈 줄은 사용자가 09-27 PPT 에서 직접 넣은 간격
      30: ['모임 관리 앱', '', '일정 · 공지 · 자료를 한곳에', '', '스마트폰 · PC · 웹 모두', '', '매주 모이는 발표 모임용'],
      31: ['전화 없이 일정 확정', '', '공지는 10초 안에 디스코드로', '', '녹화 파일은 한 번에 올림', '', '11월 20일 1차 오픈']}
S4 = {52: '일정은 전화로', 51: '한 명씩 돌려 되는 시간을 모은다',
      54: '공지는 매번 새로', 53: '같은 공지를 매주 AI 로 다시 쓴다',
      56: '녹화는 찾아서', 55: '폴더를 뒤져 올리고 링크를 보낸다',
      58: '자료는 여기저기', 57: '메신저와 드라이브에 흩어진다'}
NOW_FLOW = ['전화 돌리기', '겹치는 시간\n손으로 찾기', 'AI 로\n공지 다듬기', '메신저에\n올리기', '녹화 찾아\n올리기']
COMPARE = (['가능 시간 모으기', '공지', '자료 보관'],
           [('When2meet · Doodle', [1, 0, 0]), ('디스코드 이벤트', [0, 1, 0]), ('Meetup', [0, 1, 0]), ('주간발표회 앱', [1, 1, 1])])
S5 = {39: '지금은 다섯 단계', 37: '비슷한 서비스는 한 가지씩'}
S6 = {56: '일정 조율', 59: '안 되는 날만 고르면 투표까지',
      58: '디스코드 공지', 61: '템플릿 채우고 발송하기',
      57: '자료 올리기', 60: '녹화 폴더에서 바로 올림',
      62: '서비스 관리', 63: '역할 · 팀 · 대상별 공지'}
SCREENS = [('B-01', '홈'), ('C-03', '가능 시간 입력'), ('C-05', '후보 투표'), ('D-02', '공지 작성')]
S7 = {66: ['요구분석 · 설계', 'W1 ~ W4'], 71: ['인프라 구축', 'W1 ~ W8'], 72: ['기능 개발', 'W2 ~ W7'],
      73: ['사용자 테스트', 'W4 ~ W9'], 74: ['1차 오픈', '11월 20일']}
GOALS = [('0통', '일정 잡는 전화'), ('10초', '공지 도착'), ('1번', '녹화 올리기')]
STEPS = [('지금', 4, '전화 · 시간 찾기 · 공지 쓰기 · 올리기'), ('앱', 2, '가능 시간 입력 · 투표')]
S10 = {68: '일정 지연', 72: ['9주 · 개발 2명', '기능마다 담당 한 명, 매주 점검'],
       69: '처음 쓰는 기술', 73: ['Dart · Flutter 처음', '1주차에 골격부터 만든다'],
       70: '안 쓰면 끝', 74: ['알림을 안 보면 안 쓴다', '디스코드로 같이 알린다'],
       71: '서버 한 대', 75: ['미니PC 한 대에 전부', 'RAID 1 · 주 1회 백업']}
# 09 현황: WBS 영역별 주차(v2.5.0)와 1주차에 실제로 만든 산출물. 완료 비율은 WBS 에 아직 안 적혀 있어 쓰지 않는다.
GANTT = [('요구분석 · 설계', 1, 4), ('인프라 구축', 1, 8), ('기능 개발', 2, 7), ('사용자 테스트', 4, 9)]
NOW_WEEK, NOW_LABEL = 1, '지금 · 1주차 끝 (9/27)'
DONE = [('thumb_req.png', '요구사항서 v2.3.0'), ('thumb_sd.png', '화면설계서 v1.4.0'), ('thumb_demo.png', '클릭 데모'), ('ARCH', '구성 아키텍처 v1.2.0')]
NEXT = '2주차에는 ERD, API 설계, 서버 구성, 도메인, 프로젝트 골격을 만든다'
TEAM = [  # (사진 그룹, 이름칸, 역할칸, 설명칸, 이름, 역할, 설명)
    (30, 45, 44, 46, '홍길동', '팀장 · PM', ['인프라 · 설계', '서비스 관리자 화면']),
    (32, 41, 40, 49, '김철수', '설계 · 개발', ['프로젝트 골격', '회원 · 일정']),
    (34, 42, 43, 50, '이영희', '디자인 · QA', ['와이어프레임', '화면 검수 · 테스트']),
    (36, 47, 48, 51, '박민수', '개발', ['파일 업로드', '디스코드 공지']),
]


# ---------------------------------------------------------------- 도우미
def shape(slide, sid):
    for sh in slide.shapes:
        if sh.shape_id == sid:
            return sh
    raise KeyError(sid)


def set_text(sh, lines, size=None, bold=None, color=None, align=None):
    """첫 문단·첫 글자의 서식을 살리고 글만 바꾼다."""
    if isinstance(lines, str):
        lines = [lines]
    tf = sh.text_frame
    p0 = tf.paragraphs[0]._p
    r0 = p0.find(qn('a:r'))
    rpr = copy.deepcopy(r0.find(qn('a:rPr'))) if r0 is not None and r0.find(qn('a:rPr')) is not None else None
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    for el in list(p0):
        if el.tag not in (qn('a:pPr'),):
            p0.remove(el)
    for i, line in enumerate(lines):
        p = p0 if i == 0 else copy.deepcopy(p0)
        if i:
            for el in list(p):
                if el.tag != qn('a:pPr'):
                    p.remove(el)
            tf._txBody.append(p)
        r = p.makeelement(qn('a:r'), {})
        if rpr is not None:
            r.append(copy.deepcopy(rpr))
        t = r.makeelement(qn('a:t'), {})
        t.text = line
        r.append(t)
        p.append(r)
    for p in tf.paragraphs:
        if align is not None:
            p.alignment = align
        for r in p.runs:
            if size:
                r.font.size = Pt(size)
            if bold is not None:
                r.font.bold = bold
            if color:
                r.font.color.rgb = RGBColor.from_string(color.lstrip('#'))


def drop(slide, sid):
    el = shape(slide, sid)._element
    el.getparent().remove(el)


def textbox(slide, x, y, w, h, text, size=14, color=INK, bold=False, align=PP_ALIGN.CENTER):
    tb = slide.shapes.add_textbox(Mm(x), Mm(y), Mm(w), Mm(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(text if isinstance(text, list) else [text]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size, r.font.bold, r.font.name = Pt(size), bold, '맑은 고딕'
        r.font.color.rgb = RGBColor.from_string(color.lstrip('#'))
    return tb


def picture_fit(slide, path, x, y, w, h):
    iw, ih = Image.open(path).size
    s = min(w / iw, h / ih)
    pw, ph = iw * s, ih * s
    return slide.shapes.add_picture(str(path), Mm(x + (w - pw) / 2), Mm(y + (h - ph) / 2), Mm(pw), Mm(ph))


def all_text_shapes(shapes):
    for sh in shapes:
        if sh.shape_type == 6:
            yield from all_text_shapes(sh.shapes)
        elif sh.has_text_frame:
            yield sh


def duplicate(prs, src, keep):
    """src 슬라이드의 배경과 keep 에 든 도형만 복제한 새 슬라이드. 그림 관계(r:embed)도 옮긴다."""
    new = prs.slides.add_slide(src.slide_layout)
    for sh in list(new.shapes):
        sh._element.getparent().remove(sh._element)
    bg = src._element.cSld.find(qn('p:bg'))
    if bg is not None:
        new._element.cSld.insert(0, copy.deepcopy(bg))
    for sh in src.shapes:
        if sh.shape_id in keep:
            new.shapes._spTree.append(copy.deepcopy(sh._element))
    for el in new._element.iter():
        for attr in (qn('r:embed'), qn('r:link')):
            rid = el.get(attr)
            if rid and rid in src.part.rels:
                el.set(attr, new.part.relate_to(src.part.rels[rid].target_part, src.part.rels[rid].reltype))
    return new


def move(prs, slide, index):
    lst = prs.slides._sldIdLst
    sid = next(s for s in lst if prs.part.related_part(s.get(qn('r:id'))) is slide.part)
    lst.remove(sid)
    lst.insert(index, sid)


# ---------------------------------------------------------------- 그림
def fig_px(w_mm, h_mm, dpi=200):
    return plt.figure(figsize=(w_mm / 25.4, h_mm / 25.4), dpi=dpi)


def draw_now_flow(path, w, h):
    fig = fig_px(w, h); ax = fig.add_axes([0, 0, 1, 1]); ax.set_xlim(0, 100); ax.set_ylim(0, 100); ax.axis('off')
    n = len(NOW_FLOW); bw, gap = 15.5, 5
    x0 = (100 - (n * bw + (n - 1) * gap)) / 2
    for i, t in enumerate(NOW_FLOW):
        x = x0 + i * (bw + gap)
        ax.add_patch(FancyBboxPatch((x, 38), bw, 34, boxstyle='round,pad=0.4,rounding_size=2', fc=OLIVE if i < 2 else GRAY, ec='none'))
        ax.text(x + bw / 2, 55, t, ha='center', va='center', color='white', fontsize=17, fontweight='bold', linespacing=1.3)
        ax.text(x + bw / 2, 28, f'{i + 1}', ha='center', va='center', color=DARK, fontsize=20, fontweight='bold')
        if i < n - 1:
            ax.annotate('', xy=(x + bw + gap - 0.6, 55), xytext=(x + bw + 0.6, 55), arrowprops=dict(arrowstyle='-|>', color=DARK, lw=1.6))
    ax.text(50, 10, '다섯 단계를 모두 사람이 한다', ha='center', va='center', color=INK, fontsize=19)
    fig.savefig(path, transparent=True); plt.close(fig)


def draw_compare(path, w, h):
    cols, rows = COMPARE
    fig = fig_px(w, h); ax = fig.add_axes([0, 0, 1, 1]); ax.set_xlim(0, 100); ax.set_ylim(0, 100); ax.axis('off')
    cx = [52, 70, 88]
    for x, c in zip(cx, cols):
        ax.text(x, 90, c, ha='center', va='center', fontsize=16, color=DARK, fontweight='bold')
    for i, (name, marks) in enumerate(rows):
        y = 70 - i * 17
        mine = i == len(rows) - 1
        if mine:
            ax.add_patch(FancyBboxPatch((2, y - 7), 96, 14, boxstyle='round,pad=0.2,rounding_size=2', fc=LIGHT, ec='none'))
        ax.text(5, y, name, ha='left', va='center', fontsize=17, color=INK, fontweight='bold' if mine else 'normal')
        for x, m in zip(cx, marks):
            if m:
                ax.add_patch(Circle((x, y), 3.6, fc=DARK if mine else OLIVE, ec='none'))
            else:
                ax.plot([x - 2.5, x + 2.5], [y, y], color=GRAY, lw=2)
    fig.savefig(path, transparent=True); plt.close(fig)


def draw_steps(path, w, h):
    fig = fig_px(w, h); ax = fig.add_axes([0.14, 0.12, 0.8, 0.7])
    labels = [s[0] for s in STEPS][::-1]; vals = [s[1] for s in STEPS][::-1]; notes = [s[2] for s in STEPS][::-1]
    bars = ax.barh(labels, vals, color=[OLIVE, GRAY], height=0.55)  # 아래부터: 앱(강조), 지금
    for b, v, nt in zip(bars, vals, notes):
        ax.text(b.get_width() + 0.08, b.get_y() + b.get_height() / 2, f'{v}단계', va='center', fontsize=26, fontweight='bold', color=INK)
        ax.text(0.08, b.get_y() - 0.1, nt, va='top', fontsize=16, color=INK)
    ax.set_xlim(0, 5.2); ax.set_xticks([]); ax.tick_params(axis='y', labelsize=22, length=0)
    for sp in ax.spines.values():
        sp.set_visible(False)
    fig.text(0.05, 0.92, '일정 한 번 잡는 데 사람이 하는 일', fontsize=20, fontweight='bold', color=DARK)
    fig.text(0.05, 0.02, '설계 흐름 기준. 확정과 공지는 앱이 한다', fontsize=14, color=GRAY)
    fig.savefig(path, transparent=True); plt.close(fig)


def draw_gantt(path, w, h):
    fig = fig_px(w, h); ax = fig.add_axes([0.2, 0.2, 0.78, 0.66])
    for i, (name, a, b) in enumerate(GANTT[::-1]):
        ax.barh(i, b - a + 1, left=a - 0.5, height=0.55, color=OLIVE if a <= NOW_WEEK <= b else GRAY)
    ax.set_yticks(range(len(GANTT))); ax.set_yticklabels([g[0] for g in GANTT[::-1]], fontsize=17, color=INK)
    ax.set_xlim(0.5, 9.5); ax.set_xticks(range(1, 10)); ax.set_xticklabels([f'W{k}' for k in range(1, 10)], fontsize=14, color=INK)
    ax.axvline(NOW_WEEK + 0.5, color='#8B1A1A', lw=2.5)
    ax.text(NOW_WEEK + 0.6, len(GANTT) - 0.45, NOW_LABEL, color='#8B1A1A', fontsize=15, fontweight='bold', va='bottom')
    ax.text(9.45, len(GANTT) - 0.45, '11/20 1차 오픈', ha='right', va='bottom', fontsize=15, color=DARK, fontweight='bold')
    for sp in ('top', 'right', 'left'):
        ax.spines[sp].set_visible(False)
    ax.tick_params(length=0); ax.grid(axis='x', color=LIGHT, lw=1); ax.set_axisbelow(True)
    fig.savefig(path, transparent=True); plt.close(fig)


def capture_screens():
    out = []
    with sync_playwright() as p:
        b = p.chromium.launch()
        pg = b.new_page(viewport={'width': 1500, 'height': 1000}, device_scale_factor=2)
        for sid, _ in SCREENS:
            pg.goto(DEMO.as_uri() + f'?shot&dev=mob&t={sid}#{sid}')
            pg.evaluate('document.fonts.ready'); pg.wait_for_timeout(500)
            f = WORK / f'shot_{sid}.png'
            pg.locator('#device').screenshot(path=str(f))
            im = Image.open(f)
            im.crop((0, 0, im.width, min(im.height, im.width * 800 // 390))).save(f)
            out.append(f)
        pg.goto(DEMO.as_uri() + '?shot&dev=web&t=B-01d#B-01')
        pg.evaluate('document.fonts.ready'); pg.wait_for_timeout(500)
        pg.locator('#device').screenshot(path=str(WORK / 'thumb_demo.png'))
        b.close()
    return out


# ---------------------------------------------------------------- 만들기
def build():
    WORK.mkdir(exist_ok=True)
    OUT.parent.mkdir(exist_ok=True)
    prs = Presentation(str(TPL))
    S = list(prs.slides)
    cover, toc, s3, s4, s5, s6, s7, s8, s9, s10, s11, s12, s13 = S

    for sl in S:  # 양식 꼬리말 주소
        for sh in all_text_shapes(sl.shapes):
            if 'reallygreatsite' in sh.text_frame.text and sh.shape_id not in (28, 29):
                set_text(sh, REPO_SHORT)

    set_text(shape(cover, 19), '주간발표회')
    set_text(shape(cover, 18), '프로젝트 제안서')
    set_text(shape(cover, 20), '주간발표회')   # 사용자 09-27 수정

    nums = [18, 20, 22, 24, 26, 28, 29, 31, 33, 35]
    names = [19, 21, 37, 23, 25, 27, 30, 32, 34, 36]
    for i, (n, m) in enumerate(zip(nums, names)):
        if i < len(CONTENTS):
            set_text(shape(toc, n), f'{i + 1:02d}'); set_text(shape(toc, m), CONTENTS[i])
        else:
            drop(toc, n); drop(toc, m)

    for sid, v in S3.items():
        set_text(shape(s3, sid), v, size=26 if isinstance(v, list) else None)
    for sid, v in S4.items():
        set_text(shape(s4, sid), v, size=22 if sid in (51, 53, 55, 57) else None)

    # 03 현황 분석: 양식 그래프 자리에 흐름도와 비교표
    for pid, fn, drawer in ((19, 'now_flow.png', draw_now_flow), (18, 'compare.png', draw_compare)):
        pic = shape(s5, pid)
        x, w = {19: 31.1, 18: 253.9}[pid], Emu(pic.width).mm   # 가로 위치는 사용자 09-27 수정값
        y, h = 114, 142
        drawer(WORK / fn, w, h)
        pic._element.getparent().remove(pic._element)
        s5.shapes.add_picture(str(WORK / fn), Mm(x), Mm(y), Mm(w), Mm(h))
    for sid, v in S5.items():
        set_text(shape(s5, sid), v)
    drop(s5, 38); drop(s5, 34)

    for sid, v in S6.items():
        set_text(shape(s6, sid), v)

    # 04 화면 미리보기, 04 시스템 구성 (새 장)
    keep = {2, 9, 12, 17, 54, 55}
    sa = duplicate(prs, s6, keep)
    set_text(shape(sa, 54), '화면 미리보기')
    shots = capture_screens()
    pw, gap, top, ph = 68, 16, 78, 140
    x0 = 53 + (399 - (len(shots) * pw + (len(shots) - 1) * gap)) / 2
    for i, (f, (_, cap)) in enumerate(zip(shots, SCREENS)):
        x = x0 + i * (pw + gap)
        picture_fit(sa, f, x, top, pw, ph)
        textbox(sa, x - 8, top + ph + 2, pw + 16, 11, cap, size=22, bold=True)
    textbox(sa, 53, 244, 399, 10, '스마트폰 앱 화면. PC 앱과 웹도 같은 화면을 쓴다', size=17, color=GRAY)
    sb = duplicate(prs, s6, keep)
    set_text(shape(sb, 54), '시스템 구성')
    picture_fit(sb, ARCH, 53, 76, 399, 180)

    for sid, v in S7.items():
        set_text(shape(s7, sid), v)

    # 06 기대 효과: 게이지와 가짜 그래프를 목표 수치와 단계 막대로
    set_text(shape(s9, 19), '06')
    chart = shape(s9, 24)
    cx, cy, cw, ch = (Emu(v).mm for v in (chart.left, chart.top, chart.width, chart.height))
    for pid in (22, 23, 24):
        drop(s9, pid)
    draw_steps(WORK / 'steps.png', cw, ch)
    s9.shapes.add_picture(str(WORK / 'steps.png'), Mm(cx), Mm(cy), Mm(cw), Mm(ch))
    num_src, lab_src = shape(s9, 20), shape(s9, 25)
    boxes = [(num_src, lab_src), (shape(s9, 21), shape(s9, 26))]
    boxes.append((copy.deepcopy(num_src._element), copy.deepcopy(lab_src._element)))
    s9.shapes._spTree.append(boxes[2][0]); s9.shapes._spTree.append(boxes[2][1])
    boxes[2] = (s9.shapes[-2], s9.shapes[-1])
    for i, ((nb, lb), (big, small)) in enumerate(zip(boxes, GOALS)):
        x = 40 + i * 76
        nb.left, nb.top, nb.width, nb.height = Mm(x), Mm(115), Mm(72), Mm(42)
        lb.left, lb.top, lb.width, lb.height = Mm(x), Mm(160), Mm(72), Mm(12)
        set_text(nb, big); set_text(lb, small, size=22)

    for sid, v in S10.items():
        set_text(shape(s10, sid), v, size=22 if isinstance(v, list) else None)
    set_text(shape(s10, 67), '07')
    drop(s10, 18); drop(s10, 19)

    set_text(shape(s12, 39), '08')
    set_text(shape(s12, 38), '역할 분배')   # 사용자 09-27 수정. 목차 장은 '팀 소개' 그대로
    for i, (gid, nid, rid, bid, name, role, body) in enumerate(TEAM):
        set_text(shape(s12, nid), name); set_text(shape(s12, rid), role); set_text(shape(s12, bid), body, size=17)
        grp = shape(s12, gid)
        gx, gy, gw = Emu(grp.left).mm, Emu(grp.top).mm, Emu(grp.width).mm
        grp._element.getparent().remove(grp._element)
        ov = s12.shapes.add_shape(MSO_SHAPE.OVAL, Mm(gx), Mm(gy), Mm(gw), Mm(gw))
        ov.fill.solid(); ov.fill.fore_color.rgb = RGBColor.from_string([DARK, OLIVE, GRAY, '#6B7058'][i].lstrip('#'))
        ov.line.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); ov.line.width = Pt(4); ov.shadow.inherit = False
        tf = ov.text_frame; tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        r = tf.paragraphs[0].add_run(); r.text = name[0]
        r.font.size, r.font.bold, r.font.name = Pt(60), True, '맑은 고딕'; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    sc = duplicate(prs, s6, keep)
    set_text(shape(sc, 54), '현황'); set_text(shape(sc, 55), '09')
    draw_gantt(WORK / 'gantt.png', 399, 78)
    sc.shapes.add_picture(str(WORK / 'gantt.png'), Mm(53), Mm(74), Mm(399), Mm(78))
    textbox(sc, 53, 154, 150, 10, '1주차에 만든 것', size=19, color=DARK, bold=True, align=PP_ALIGN.LEFT)
    tw, tg = 90, 13
    for i, (fn, cap) in enumerate(DONE):
        x = 53 + i * (tw + tg)
        pic = picture_fit(sc, ARCH if fn == 'ARCH' else WORK / fn, x, 166, tw, 62)
        pic.line.color.rgb = RGBColor(0xC8, 0xC8, 0xC2); pic.line.width = Pt(1)
        textbox(sc, x - 4, 229, tw + 8, 10, cap, size=16, bold=True)
    textbox(sc, 53, 243, 399, 10, NEXT, size=17, color=GRAY)

    set_text(shape(s13, 33), '주간발표회')
    set_text(shape(s13, 28), REPO, size=15)
    for sid in (27, 29, 24, 25):
        drop(s13, sid)

    # 예산 장을 빼고, 새 장 둘을 제안 솔루션 뒤로
    lst = prs.slides._sldIdLst
    for gone in (s8, s11):
        sid = next(x for x in lst if prs.part.related_part(x.get(qn('r:id'))) is gone.part)
        prs.part.drop_rel(sid.get(qn('r:id'))); lst.remove(sid)
    move(prs, sa, 6); move(prs, sb, 7); move(prs, sc, len(lst) - 2)

    for sl in prs.slides:  # 복제로 겹친 도형 id 를 다시 매긴다
        for k, el in enumerate(sl.shapes._spTree.iter(qn('p:cNvPr')), start=2):
            el.set('id', str(k))
    prs.save(str(OUT))
    print('saved', OUT, len(prs.slides), '장')


if __name__ == '__main__':
    build()
